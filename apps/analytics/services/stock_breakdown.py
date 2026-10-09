import io
import datetime
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any
from collections import defaultdict
from dateutil.relativedelta import relativedelta
from django.utils import timezone
from django.db.models import QuerySet, Sum, Q
from django.db.models.functions import Coalesce
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from apps.products.models import Stock, ProductCategory, ProductClass, Product
from apps.sales.models import Warehouse


@dataclass
class StockBreakdownService:
    queryset: QuerySet
    dimension: str = 'productcategory_productclass_product'
    user: Any | None = None
    cleaned_data: dict[str, Any] | None = None

    dimension_config: dict[str, Any] = field(init=False)
    active_warehouses: list[Warehouse] = field(default_factory=list, init=False)
    warehouse_ids: list[str] = field(default_factory=list, init=False)

    DIMENSION_CONFIG = {
        'productcategory_productclass_product': {
            'label': 'Categoría de producto → Clase de producto → Producto → Lote',
            'depth': 4,
            'l1_id': 'product__product_class__product_category_id',
            'l1_model': ProductCategory,
            'l1_label': 'Categoría de producto',
            'l2_id': 'product__product_class_id',
            'l2_model': ProductClass,
            'l2_label': 'Clase de producto',
            'l3_id': 'product_id',
            'l3_model': Product,
            'l3_label': 'Producto',
            'l4_id': 'lot_number',
            'l4_label': 'Lote',
        },
    }

    def __post_init__(self):
        self.queryset = self.queryset.order_by()
        if self.dimension not in self.DIMENSION_CONFIG:
            self.dimension = 'productcategory_productclass_product'
        self.dimension_config = self.DIMENSION_CONFIG[self.dimension]

        self._resolve_warehouses()
        self.queryset = self.queryset.filter(warehouse_id__in=self.warehouse_ids)

    def _resolve_warehouses(self) -> None:
        """
        resolves active warehouses based on cleaned_data or default type WAREHOUSE
        """
        selected_warehouses = (self.cleaned_data or {}).get('warehouse')
        if selected_warehouses:
            self.active_warehouses = list(
                Warehouse.objects.filter(id__in=[w.pk if hasattr(w, 'pk') else w for w in selected_warehouses])
                .order_by('name')
            )
        else:
            self.active_warehouses = list(
                Warehouse.objects.filter(warehouse_type=Warehouse.WarehouseTypeChoices.WAREHOUSE)
                .order_by('name')
            )

        self.warehouse_ids = [w.id for w in self.active_warehouses]

    @property
    def l1_id_field(self) -> str:
        return self.dimension_config['l1_id']

    def _init_warehouse_totals(self) -> dict[str, float]:
        return {w_id: 0.0 for w_id in self.warehouse_ids}

    def _flatten_warehouse_totals(self, totals_dict: dict[str, float]) -> tuple[list[dict[str, Any]], float]:
        result = []
        overall = 0.0
        for w in self.active_warehouses:
            qty = totals_dict.get(w.id, 0.0)
            overall += qty
            result.append({
                'warehouse_id': w.id,
                'warehouse_name': w.name,
                'quantity': round(qty, 2),
            })
        return result, round(overall, 2)

    def get_level_1_queryset(self) -> QuerySet:
        """
        returns queryset of distinct level 1 ids ordered by overall total quantity
        """
        l1_id = self.l1_id_field
        return (
            self.queryset.values(l1_id)
            .annotate(total_overall=Sum('quantity'))
            .filter(total_overall__gt=0)
            .order_by('-total_overall')
        )

    def get_level_1_items(self, top_l1_ids: list[Any]) -> list[dict[str, Any]]:
        """
        fetches and returns level 1 items with their warehouse quantities for the initial page load
        """
        if not top_l1_ids:
            return []

        l1_id = self.l1_id_field
        l1_model = self.dimension_config['l1_model']
        depth = self.dimension_config['depth']

        data_list = (
            self.queryset
            .filter(**{f'{l1_id}__in': top_l1_ids})
            .values(l1_id, 'warehouse_id')
            .annotate(total_quantity=Sum('quantity'))
            .order_by()
        )

        l1_names = dict(
            l1_model.objects.filter(id__in=top_l1_ids).values_list('id', 'name')
        )

        totals_map: dict[Any, dict[str, float]] = {
            cid: self._init_warehouse_totals() for cid in top_l1_ids
        }

        for row in data_list:
            k1_id = row.get(l1_id)
            w_id = row.get('warehouse_id')
            qty = float(row.get('total_quantity') or 0.0)

            if w_id in totals_map.get(k1_id, {}):
                totals_map[k1_id][w_id] += qty

        result = []
        for k1_id in top_l1_ids:
            name = l1_names.get(k1_id) or 'Sin registro'
            node_id = f"n1_{k1_id}"
            warehouse_totals, total_overall = self._flatten_warehouse_totals(totals_map[k1_id])

            if total_overall <= 0:
                continue

            result.append({
                'id': k1_id,
                'name': name,
                'level': 1,
                'next_level': 2,
                'depth': depth,
                'has_children': depth > 1,
                'node_id': node_id,
                'l1_id': k1_id,
                'totals': warehouse_totals,
                'total_overall': total_overall,
            })

        return result

    def _get_level_4_lots(self, parent_filters: dict[str, Any]) -> list[dict[str, Any]]:
        """
        fetches and returns level 4 child items representing lots for a given product
        """
        depth = self.dimension_config['depth']
        product_id = parent_filters.get('l3_id')
        if not product_id:
            return []

        stock_filters = {'product_id': product_id}
        if parent_filters.get('l1_id') and str(parent_filters['l1_id']) != 'otros':
            stock_filters[self.dimension_config['l1_id']] = parent_filters['l1_id']
        if parent_filters.get('l2_id') and str(parent_filters['l2_id']) != 'otros':
            stock_filters[self.dimension_config['l2_id']] = parent_filters['l2_id']

        data_list = (
            self.queryset
            .filter(**stock_filters)
            .values('lot_number', 'expiration_date', 'warehouse_id')
            .annotate(total_quantity=Sum('quantity'))
            .order_by()
        )

        child_totals: dict[tuple[str, Any], dict[str, float]] = defaultdict(self._init_warehouse_totals)
        child_qty_sum: dict[tuple[str, Any], float] = defaultdict(float)

        for row in data_list:
            lot_num = (row.get('lot_number') or '').strip()
            exp_date = row.get('expiration_date')
            w_id = row.get('warehouse_id')
            qty = float(row.get('total_quantity') or 0.0)

            if w_id in self.warehouse_ids:
                key = (lot_num, exp_date)
                child_totals[key][w_id] += qty
                child_qty_sum[key] += qty

        #include lots with quantity != 0
        active_lots = [(k, tot) for k, tot in child_qty_sum.items() if round(tot, 2) != 0]

        #fefo order earliest expiration date first, no expiration date last, then descending by quantity
        def _lot_sort_key(item):
            (lot_num, exp_date), tot = item
            has_date = 0 if exp_date is not None else 1
            return (has_date, exp_date or datetime.date.max, -tot, lot_num)

        sorted_lots = sorted(active_lots, key=_lot_sort_key)

        parent_node_id = parent_filters.get('parent_node_id', '')

        def _build_ancestor_classes(p_node_id: str) -> str:
            if not p_node_id:
                return ''
            parts = p_node_id.split('_')
            classes = []
            for i in range(2, len(parts) + 1):
                prefix = '_'.join(parts[:i])
                classes.append(f'child-{prefix}')
            return ' '.join(classes)

        ancestor_classes = _build_ancestor_classes(parent_node_id)
        today = timezone.localdate()
        result = []

        for idx, ((lot_num, exp_date), _) in enumerate(sorted_lots, start=1):
            key = (lot_num, exp_date)
            warehouse_totals, total_overall = self._flatten_warehouse_totals(child_totals[key])
            safe_lot = re.sub(r'[^a-zA-Z0-9]', '_', str(lot_num)) if lot_num else 'sin_lote'
            exp_str = exp_date.strftime('%Y%m%d') if exp_date else 'noexp'
            node_id = f"{parent_node_id}_l_{safe_lot}_{exp_str}_{idx}"

            is_expired = bool(exp_date and exp_date <= today)

            result.append({
                'id': f"{lot_num}_{exp_str}",
                'name': lot_num if lot_num else 'Sin lote',
                'lot_number': lot_num,
                'expiration_date': exp_date,
                'is_expired': is_expired,
                'is_lot': True,
                'is_product': False,
                'level': 4,
                'next_level': 5,
                'depth': depth,
                'has_children': False,
                'node_id': node_id,
                'parent_node_id': parent_node_id,
                'ancestor_classes': ancestor_classes,
                'l1_id': parent_filters.get('l1_id'),
                'l2_id': parent_filters.get('l2_id'),
                'l3_id': parent_filters.get('l3_id'),
                'totals': warehouse_totals,
                'total_overall': total_overall,
            })

        return result

    def get_level_children(self, target_level: int, parent_filters: dict[str, Any]) -> list[dict[str, Any]]:
        """
        fetches and returns child items for a given level and parent filter criteria
        """
        depth = self.dimension_config['depth']
        if target_level > depth:
            return []

        if target_level == 4:
            return self._get_level_4_lots(parent_filters)

        target_id_field = self.dimension_config[f'l{target_level}_id']
        target_model = self.dimension_config[f'l{target_level}_model']

        # build filter criteria for stock from parent identifiers
        stock_filters = {}
        for lvl in range(1, target_level):
            parent_val = parent_filters.get(f'l{lvl}_id')
            if parent_val is not None and str(parent_val).strip() != '' and str(parent_val) != 'otros':
                field_name = self.dimension_config[f'l{lvl}_id']
                stock_filters[field_name] = parent_val

        data_list = (
            self.queryset
            .filter(**stock_filters)
            .values(target_id_field, 'warehouse_id')
            .annotate(total_quantity=Sum('quantity'))
            .order_by()
        )

        child_totals: dict[Any, dict[str, float]] = defaultdict(self._init_warehouse_totals)
        child_qty_sum: dict[Any, float] = defaultdict(float)

        for row in data_list:
            c_id = row.get(target_id_field)
            if c_id is None:
                continue
            w_id = row.get('warehouse_id')
            qty = float(row.get('total_quantity') or 0.0)

            if w_id in self.warehouse_ids:
                child_totals[c_id][w_id] += qty
                child_qty_sum[c_id] += qty

        # filter out children with 0 stock and sort descending by total quantity
        active_children = [(cid, tot) for cid, tot in child_qty_sum.items() if round(tot, 2) != 0]
        sorted_children = sorted(active_children, key=lambda x: x[1], reverse=True)

        all_ids = [c_id for c_id, _ in sorted_children]
        target_names = dict(
            target_model.objects.filter(id__in=all_ids).values_list('id', 'name')
        )

        parent_node_id = parent_filters.get('parent_node_id', '')

        def _build_ancestor_classes(p_node_id: str) -> str:
            if not p_node_id:
                return ''
            parts = p_node_id.split('_')
            classes = []
            for i in range(2, len(parts) + 1):
                prefix = '_'.join(parts[:i])
                classes.append(f'child-{prefix}')
            return ' '.join(classes)

        ancestor_classes = _build_ancestor_classes(parent_node_id)
        result = []
        is_product = (target_level == 3)

        for c_id, _ in sorted_children:
            name = target_names.get(c_id) or 'Sin registro'
            node_id = f"{parent_node_id}_{c_id}" if parent_node_id else f"n{target_level}_{c_id}"
            warehouse_totals, total_overall = self._flatten_warehouse_totals(child_totals[c_id])

            item_data = {
                'id': c_id,
                'name': name,
                'is_product': is_product,
                'is_lot': False,
                'level': target_level,
                'next_level': target_level + 1,
                'depth': depth,
                'has_children': target_level < depth,
                'node_id': node_id,
                'parent_node_id': parent_node_id,
                'ancestor_classes': ancestor_classes,
                'l1_id': parent_filters.get('l1_id'),
                'totals': warehouse_totals,
                'total_overall': total_overall,
            }
            if target_level >= 2:
                item_data['l2_id'] = c_id if target_level == 2 else parent_filters.get('l2_id')
            if target_level >= 3:
                item_data['l3_id'] = c_id if target_level == 3 else parent_filters.get('l3_id')

            result.append(item_data)

        return result

    def get_kpis(self) -> dict[str, Any]:
        today = timezone.localdate()
        m0_start = today.replace(day=1)
        start_1 = m0_start
        end_1 = (m0_start + relativedelta(months=2)) - relativedelta(days=1)
        start_2 = m0_start + relativedelta(months=2)
        end_2 = (m0_start + relativedelta(months=4)) - relativedelta(days=1)
        start_3 = m0_start + relativedelta(months=4)
        end_3 = (m0_start + relativedelta(months=6)) - relativedelta(days=1)
        start_4 = m0_start + relativedelta(months=6)

        #deduplicate stocks by primary key
        unique_stocks = Stock.objects.filter(id__in=self.queryset.values('id'))

        agg = unique_stocks.aggregate(
            total_stock=Coalesce(Sum('quantity'), Decimal('0.00')),
            expired_units=Coalesce(
                Sum('quantity', filter=Q(expiration_date__isnull=False, expiration_date__lte=today)),
                Decimal('0.00'),
            ),
            expires_0_2_months=Coalesce(
                Sum('quantity', filter=Q(expiration_date__isnull=False, expiration_date__gte=start_1, expiration_date__lte=end_1)),
                Decimal('0.00'),
            ),
            expires_3_4_months=Coalesce(
                Sum('quantity', filter=Q(expiration_date__isnull=False, expiration_date__gte=start_2, expiration_date__lte=end_2)),
                Decimal('0.00'),
            ),
            expires_5_6_months=Coalesce(
                Sum('quantity', filter=Q(expiration_date__isnull=False, expiration_date__gte=start_3, expiration_date__lte=end_3)),
                Decimal('0.00'),
            ),
            expires_6_plus_months=Coalesce(
                Sum('quantity', filter=Q(expiration_date__isnull=False, expiration_date__gte=start_4)),
                Decimal('0.00'),
            ),
        )

        products_with_stock = (
            unique_stocks
            .filter(product_id__isnull=False)
            .values('product_id')
            .annotate(total_qty=Sum('quantity'))
            .exclude(total_qty=0)
            .count()
        )

        expired_products = (
            unique_stocks
            .filter(product_id__isnull=False, expiration_date__isnull=False, expiration_date__lte=today)
            .values('product_id')
            .annotate(total_qty=Sum('quantity'))
            .exclude(total_qty=0)
            .count()
        )

        return {
            'products_with_stock': products_with_stock,
            'total_stock': float(agg['total_stock']),
            'expired_units': float(agg['expired_units']),
            'expired_products': expired_products,
            'expires_0_2_months': float(agg['expires_0_2_months']),
            'expires_3_4_months': float(agg['expires_3_4_months']),
            'expires_5_6_months': float(agg['expires_5_6_months']),
            'expires_6_plus_months': float(agg['expires_6_plus_months']),
        }


@dataclass
class StockBreakdownExports:
    breakdown_service: StockBreakdownService

    def export_stock_breakdown_excel(self) -> io.BytesIO:
        """
        exports the full stock breakdown hierarchical tree to an Excel workbook with openpyxl.
        creates structured tree rows with distinct styling and visual indentation for:
        Category -> Class -> Product.
        """
        service = self.breakdown_service
        warehouses = service.active_warehouses
        wh_ids = service.warehouse_ids

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Existencias por Almacén"

        FONT_FAMILY = "Segoe UI"
        HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # slate-800
        HEADER_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
        HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

        CAT_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")  # slate-200
        CAT_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color="0F172A")

        CLASS_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # slate-100
        CLASS_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="334155")

        PROD_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        PROD_FONT = Font(name=FONT_FAMILY, size=9, bold=False, color="1E293B")

        LOT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")  # slate-50
        LOT_FONT = Font(name=FONT_FAMILY, size=8.5, italic=False, color="475569")

        TOTAL_COL_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # amber-100
        TOTAL_COL_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="92400E")

        BORDER_THIN = Border(
            left=Side(style='thin', color="CBD5E1"),
            right=Side(style='thin', color="CBD5E1"),
            top=Side(style='thin', color="CBD5E1"),
            bottom=Side(style='thin', color="CBD5E1"),
        )

        NUM_FORMAT = "#,##0.00"

        headers = ["Existencias"]
        for w in warehouses:
            headers.append(f"{w.name.title()} ({w.get_warehouse_type_display().title()})")
        headers.append("Total Existencias")

        ws.append(headers)
        ws.row_dimensions[1].height = 28

        for col_idx, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGNMENT
            cell.border = BORDER_THIN

        raw_data = (
            service.queryset
            .values(
                'product__product_class__product_category_id',
                'product__product_class__product_category__name',
                'product__product_class_id',
                'product__product_class__name',
                'product_id',
                'product__name',
                'lot_number',
                'expiration_date',
                'warehouse_id',
            )
            .annotate(total_quantity=Sum('quantity'))
            .filter(total_quantity__gt=0)
            .order_by(
                'product__product_class__product_category__name',
                'product__product_class__name',
                'product__name',
                'expiration_date',
                'lot_number',
            )
        )

        tree: dict[str, Any] = {}

        for row in raw_data:
            cat_id = row['product__product_class__product_category_id'] or 'sin_categoria'
            cat_name = row['product__product_class__product_category__name'] or 'Sin Categoría'

            class_id = row['product__product_class_id'] or 'sin_clase'
            class_name = row['product__product_class__name'] or 'Sin Clase'

            prod_id = row['product_id']
            prod_name = row['product__name'] or 'Sin Nombre'

            lot_num = (row.get('lot_number') or '').strip()
            exp_date = row.get('expiration_date')
            lot_key = (lot_num, exp_date)

            w_id = row['warehouse_id']
            qty = float(row['total_quantity'] or 0.0)

            if cat_id not in tree:
                tree[cat_id] = {
                    'name': cat_name,
                    'totals': service._init_warehouse_totals(),
                    'classes': {},
                }
            tree[cat_id]['totals'][w_id] += qty

            cat_classes = tree[cat_id]['classes']
            if class_id not in cat_classes:
                cat_classes[class_id] = {
                    'name': class_name,
                    'totals': service._init_warehouse_totals(),
                    'products': {},
                }
            cat_classes[class_id]['totals'][w_id] += qty

            class_prods = cat_classes[class_id]['products']
            if prod_id not in class_prods:
                class_prods[prod_id] = {
                    'name': prod_name,
                    'totals': service._init_warehouse_totals(),
                    'lots': {},
                }
            class_prods[prod_id]['totals'][w_id] += qty

            prod_lots = class_prods[prod_id]['lots']
            if lot_key not in prod_lots:
                prod_lots[lot_key] = {
                    'lot_number': lot_num,
                    'expiration_date': exp_date,
                    'totals': service._init_warehouse_totals(),
                }
            prod_lots[lot_key]['totals'][w_id] += qty

        current_row = 2
        total_col_idx = len(headers)

        for cat_id, cat_data in tree.items():
            cat_tot = sum(cat_data['totals'].values())
            if cat_tot <= 0:
                continue

            ws.cell(row=current_row, column=1, value=f"{cat_data['name'].upper()}")
            ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="left", vertical="center")

            for idx, wid in enumerate(wh_ids, start=2):
                q = cat_data['totals'].get(wid, 0.0)
                c = ws.cell(row=current_row, column=idx, value=q)
                c.number_format = NUM_FORMAT
                c.alignment = Alignment(horizontal="right", vertical="center")

            tot_cell = ws.cell(row=current_row, column=total_col_idx, value=cat_tot)
            tot_cell.number_format = NUM_FORMAT
            tot_cell.alignment = Alignment(horizontal="right", vertical="center")

            for col_idx in range(1, total_col_idx + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.fill = CAT_FILL
                cell.font = CAT_FONT
                cell.border = BORDER_THIN
            ws.row_dimensions[current_row].height = 22
            current_row += 1

            for class_id, class_data in cat_data['classes'].items():
                class_tot = sum(class_data['totals'].values())
                if class_tot <= 0:
                    continue

                ws.cell(row=current_row, column=1, value=f"  • {class_data['name'].title()}")
                ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="left", vertical="center")

                for idx, wid in enumerate(wh_ids, start=2):
                    q = class_data['totals'].get(wid, 0.0)
                    c = ws.cell(row=current_row, column=idx, value=q)
                    c.number_format = NUM_FORMAT
                    c.alignment = Alignment(horizontal="right", vertical="center")

                tot_cell = ws.cell(row=current_row, column=total_col_idx, value=class_tot)
                tot_cell.number_format = NUM_FORMAT
                tot_cell.alignment = Alignment(horizontal="right", vertical="center")

                for col_idx in range(1, total_col_idx + 1):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.fill = CLASS_FILL
                    cell.font = CLASS_FONT
                    cell.border = BORDER_THIN
                ws.row_dimensions[current_row].height = 20
                current_row += 1

                for prod_id, prod_data in class_data['products'].items():
                    prod_tot = sum(prod_data['totals'].values())
                    if prod_tot <= 0:
                        continue

                    ws.cell(row=current_row, column=1, value=f"      {prod_id.upper()} - {prod_data['name'].title()}")
                    ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="left", vertical="center")

                    for idx, wid in enumerate(wh_ids, start=2):
                        q = prod_data['totals'].get(wid, 0.0)
                        c = ws.cell(row=current_row, column=idx, value=q if q > 0 else 0.0)
                        c.number_format = NUM_FORMAT
                        c.alignment = Alignment(horizontal="right", vertical="center")

                    tot_cell = ws.cell(row=current_row, column=total_col_idx, value=prod_tot)
                    tot_cell.number_format = NUM_FORMAT
                    tot_cell.alignment = Alignment(horizontal="right", vertical="center")

                    for col_idx in range(1, total_col_idx):
                        cell = ws.cell(row=current_row, column=col_idx)
                        cell.fill = PROD_FILL
                        cell.font = PROD_FONT
                        cell.border = BORDER_THIN

                    tot_cell.fill = TOTAL_COL_FILL
                    tot_cell.font = TOTAL_COL_FONT
                    tot_cell.border = BORDER_THIN

                    ws.row_dimensions[current_row].height = 18
                    current_row += 1

                    for (lot_num, exp_date), lot_data in prod_data.get('lots', {}).items():
                        lot_tot = sum(lot_data['totals'].values())
                        if lot_tot <= 0:
                            continue

                        lot_label = lot_num if lot_num else "Sin lote"
                        exp_str = f" (Cad: {exp_date:%d/%m/%Y})" if exp_date else " (Sin caducidad)"

                        ws.cell(row=current_row, column=1, value=f"          └ Lote: {lot_label}{exp_str}")
                        ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="left", vertical="center")

                        for idx, wid in enumerate(wh_ids, start=2):
                            q = lot_data['totals'].get(wid, 0.0)
                            c = ws.cell(row=current_row, column=idx, value=q if q > 0 else 0.0)
                            c.number_format = NUM_FORMAT
                            c.alignment = Alignment(horizontal="right", vertical="center")

                        lot_tot_cell = ws.cell(row=current_row, column=total_col_idx, value=lot_tot)
                        lot_tot_cell.number_format = NUM_FORMAT
                        lot_tot_cell.alignment = Alignment(horizontal="right", vertical="center")

                        for col_idx in range(1, total_col_idx + 1):
                            cell = ws.cell(row=current_row, column=col_idx)
                            cell.fill = LOT_FILL
                            cell.font = LOT_FONT
                            cell.border = BORDER_THIN

                        ws.row_dimensions[current_row].height = 17
                        current_row += 1

        ws.freeze_panes = "B2"

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)

        ws.column_dimensions['A'].width = 46

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
