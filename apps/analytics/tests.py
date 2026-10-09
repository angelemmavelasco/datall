'''
tests made with ia to make sure the code is correct and all environment works correctly
'''


import csv
import io
from datetime import date, timedelta
from decimal import Decimal
import openpyxl
from django.contrib.auth import get_user_model
from django.db.models import Q, Sum
from django.template.loader import render_to_string
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from apps.analytics.filters import YearlySaleBreakdownFilter, StockBreakdownFilter
from apps.analytics.services.yearly_sale_breakdown import (
    YearlySaleBreakdownExports,
    YearlySaleBreakdownService,
)
from apps.analytics.services.stock_breakdown import (
    StockBreakdownService,
    StockBreakdownExports,
)
from apps.analytics.services.customer_kpis import (
    CustomerKpisService,
    CustomerProfileService,
    CustomerKpisExports,
)
from apps.customers.models import Customer, CustomerAssignment, CustomerType, CustomerAgreement, CommercialBenefit, AccountsReceivable
from apps.human_resources.models import BusinessUnit
from apps.products.models import Product, ProductClass, ProductCategory, Stock
from apps.sales.models import Route, RouteType, SaleChannel, SaleTransaction, Warehouse
from apps.sales.services.sale_transactions import SaleTransactionsService

User = get_user_model()


class YearlySaleBreakdownAggregatesTestCase(TestCase):
    """
    Test suite validating YearlySaleBreakdownService, YearlySaleBreakdownFilter,
    and YearlySaleBreakdownExports by comparing system outputs directly against
    raw database aggregate queries (Sum('net_amount'), Sum('profit')).
    """

    @classmethod
    def setUpTestData(cls):
        
        cls.user = User.objects.create_user(
            username='analytics_admin',
            email='admin@datall.local',
            is_superuser=True,
            is_staff=True,
        )

        
        cls.region_north = BusinessUnit.objects.create(
            id='reg_norte',
            name='Región Norte',
            business_unit_type=BusinessUnit.BusinessUnitTypeChoices.REGION,
        )
        cls.region_south = BusinessUnit.objects.create(
            id='reg_sur',
            name='Región Sur',
            business_unit_type=BusinessUnit.BusinessUnitTypeChoices.REGION,
        )

        cls.bu_1 = BusinessUnit.objects.create(
            id='bu_monterrey',
            name='Gerencia Monterrey',
            business_unit_type=BusinessUnit.BusinessUnitTypeChoices.UNIT,
            parent=cls.region_north,
        )
        cls.bu_2 = BusinessUnit.objects.create(
            id='bu_guadalajara',
            name='Gerencia Guadalajara',
            business_unit_type=BusinessUnit.BusinessUnitTypeChoices.UNIT,
            parent=cls.region_south,
        )

        
        cls.route_type = RouteType.objects.create(id='rt_std', name='Preventa')
        cls.sale_channel = SaleChannel.objects.create(id='sc_trad', name='Tradicional')
        cls.warehouse = Warehouse.objects.create(id='wh_main', name='Almacén Central')

        cls.route_a = Route.objects.create(
            id='R-101',
            name='Ruta 101',
            business_unit=cls.bu_1,
            route_type=cls.route_type,
            sale_channel=cls.sale_channel,
            is_active=True,
        )
        cls.route_b = Route.objects.create(
            id='R-102',
            name='Ruta 102',
            business_unit=cls.bu_2,
            route_type=cls.route_type,
            sale_channel=cls.sale_channel,
            is_active=True,
        )

        
        cls.category = ProductCategory.objects.create(id='cat_farma', name='Farmacéuticos')
        cls.class_analgesic = ProductClass.objects.create(
            id='pc_analg', name='Analgésicos', product_category=cls.category
        )
        cls.class_antibiotic = ProductClass.objects.create(
            id='pc_antib', name='Antibióticos', product_category=cls.category
        )

        cls.prod_paracetamol = Product.objects.create(
            id='P-001', name='Paracetamol 500mg', product_class=cls.class_analgesic
        )
        cls.prod_ibuprofen = Product.objects.create(
            id='P-002', name='Ibuprofeno 400mg', product_class=cls.class_analgesic
        )
        cls.prod_amoxicillin = Product.objects.create(
            id='P-003', name='Amoxicilina 500mg', product_class=cls.class_antibiotic
        )

        
        cls.customer_type = CustomerType.objects.create(id='ct_ind', name='Independiente')

        cls.customer_1 = Customer.objects.create(
            id='C-001',
            name='Farmacia San José',
            registration_date=date(2023, 1, 1),
            customer_type=cls.customer_type,
        )
        cls.customer_2 = Customer.objects.create(
            id='C-002',
            name='Farmacia El Ahorro',
            registration_date=date(2023, 1, 1),
            customer_type=cls.customer_type,
        )
        cls.customer_3 = Customer.objects.create(
            id='C-003',
            name='Farmacia Moderna',
            registration_date=date(2023, 1, 1),
            customer_type=cls.customer_type,
        )

        
        CustomerAssignment.objects.create(
            customer=cls.customer_1,
            route=cls.route_b,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 12, 31),
        )
        CustomerAssignment.objects.create(
            customer=cls.customer_1,
            route=cls.route_a,
            start_date=date(2025, 1, 1),
            end_date=None,  
        )

        
        CustomerAssignment.objects.create(
            customer=cls.customer_2,
            route=cls.route_a,
            start_date=date(2024, 1, 1),
            end_date=date(2024, 12, 31),
        )

        
        CustomerAssignment.objects.create(
            customer=cls.customer_3,
            route=cls.route_b,
            start_date=date(2024, 1, 1),
            end_date=None,  
        )

        
        
        cls.tx1 = SaleTransaction.objects.create(
            doc_id='FAC-1001',
            sale_date=date(2024, 3, 15),
            customer=cls.customer_1,
            route=cls.route_b,
            warehouse=cls.warehouse,
            product_class=cls.class_analgesic,
            product=cls.prod_paracetamol,
            net_amount=Decimal('10000.00'),
            profit=Decimal('2000.00'),
            cost=Decimal('8000.00'),
            gross_amount=Decimal('11600.00'),
            quantity=Decimal('100.00'),
        )
        cls.tx2 = SaleTransaction.objects.create(
            doc_id='FAC-1002',
            sale_date=date(2024, 8, 20),
            customer=cls.customer_1,
            route=cls.route_b,
            warehouse=cls.warehouse,
            product_class=cls.class_antibiotic,
            product=cls.prod_amoxicillin,
            net_amount=Decimal('6000.00'),
            profit=Decimal('1500.00'),
            cost=Decimal('4500.00'),
            gross_amount=Decimal('6960.00'),
            quantity=Decimal('50.00'),
        )

        
        cls.tx3 = SaleTransaction.objects.create(
            doc_id='FAC-1003',
            sale_date=date(2025, 2, 10),
            customer=cls.customer_1,
            route=cls.route_a,
            warehouse=cls.warehouse,
            product_class=cls.class_analgesic,
            product=cls.prod_ibuprofen,
            net_amount=Decimal('15000.00'),
            profit=Decimal('3500.00'),
            cost=Decimal('11500.00'),
            gross_amount=Decimal('17400.00'),
            quantity=Decimal('120.00'),
        )

        
        cls.tx4 = SaleTransaction.objects.create(
            doc_id='FAC-2001',
            sale_date=date(2024, 5, 12),
            customer=cls.customer_2,
            route=cls.route_a,
            warehouse=cls.warehouse,
            product_class=cls.class_analgesic,
            product=cls.prod_paracetamol,
            net_amount=Decimal('5000.00'),
            profit=Decimal('1000.00'),
            cost=Decimal('4000.00'),
            gross_amount=Decimal('5800.00'),
            quantity=Decimal('50.00'),
        )

        
        cls.tx5 = SaleTransaction.objects.create(
            doc_id='FAC-3001',
            sale_date=date(2025, 4, 18),
            customer=cls.customer_3,
            route=cls.route_b,
            warehouse=cls.warehouse,
            product_class=cls.class_antibiotic,
            product=cls.prod_amoxicillin,
            net_amount=Decimal('8000.00'),
            profit=Decimal('1600.00'),
            cost=Decimal('6400.00'),
            gross_amount=Decimal('9280.00'),
            quantity=Decimal('80.00'),
        )

    def test_customer_dimension_aggregates_vs_db_raw(self):
        """
        Validates that filtering by Route A in customer_productclass_product:
        1. Returns only actively assigned customers (Customer 1).
        2. Customer 1 shows full historical sales (2024 + 2025), matching direct DB aggregates.
        3. Excludes Customer 2 (assignment ended) and Customer 3 (assigned to Route B).
        """
        today = timezone.localdate()
        tx_service = SaleTransactionsService(user=self.user)
        base_tx_qs = tx_service.read_transactions_by_allowed_customers()

        filter_set = YearlySaleBreakdownFilter(
            data={'dimension': 'customer_productclass_product', 'route': [self.route_a.id]},
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_set.is_valid(), filter_set.errors)
        filtered_qs = filter_set.qs

        service = YearlySaleBreakdownService(
            queryset=filtered_qs,
            dimension='customer_productclass_product',
            user=self.user,
            cleaned_data=filter_set.form.cleaned_data,
        )

        
        active_customer_ids = list(
            CustomerAssignment.objects.filter(route=self.route_a)
            .filter(Q(end_date__isnull=True) | Q(end_date__gte=today))
            .values_list('customer_id', flat=True)
        )
        self.assertEqual(active_customer_ids, [self.customer_1.id])

        db_aggregates = SaleTransaction.objects.filter(
            customer_id__in=active_customer_ids
        ).aggregate(
            net_total=Sum('net_amount'),
            profit_total=Sum('profit'),
        )

        
        l1_qs = service.get_level_1_queryset()
        l1_ids = [item['customer_id'] for item in l1_qs]
        self.assertEqual(l1_ids, [self.customer_1.id])

        items = service.get_level_1_items(l1_ids)
        self.assertEqual(len(items), 1)

        c1_item = items[0]
        self.assertEqual(c1_item['id'], self.customer_1.id)

        totals_by_year = {t['year']: t for t in c1_item['totals']}

        
        for year in [2024, 2025]:
            db_year = SaleTransaction.objects.filter(
                customer=self.customer_1, sale_date__year=year
            ).aggregate(net=Sum('net_amount'), profit=Sum('profit'))
            expected_net = float(db_year['net'] or 0.0)
            expected_profit = float(db_year['profit'] or 0.0)

            actual_net = totals_by_year[year]['net']
            actual_profit = totals_by_year[year]['profit']

            self.assertAlmostEqual(actual_net, expected_net, places=2)
            self.assertAlmostEqual(actual_profit, expected_profit, places=2)

        
        total_service_net = sum(t['net'] for t in c1_item['totals'])
        total_service_profit = sum(t['profit'] for t in c1_item['totals'])
        self.assertAlmostEqual(total_service_net, float(db_aggregates['net_total']), places=2)
        self.assertAlmostEqual(total_service_profit, float(db_aggregates['profit_total']), places=2)

    def test_level_children_aggregates_vs_db_raw(self):
        """
        Validates Level 2 (Product Class) and Level 3 (Product) child breakdowns
        against raw database aggregates for Customer 1 under Route A.
        """
        tx_service = SaleTransactionsService(user=self.user)
        base_tx_qs = tx_service.read_transactions_by_allowed_customers()

        filter_set = YearlySaleBreakdownFilter(
            data={'dimension': 'customer_productclass_product', 'route': [self.route_a.id]},
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_set.is_valid(), filter_set.errors)

        service = YearlySaleBreakdownService(
            queryset=filter_set.qs,
            dimension='customer_productclass_product',
            user=self.user,
            cleaned_data=filter_set.form.cleaned_data,
        )

        
        l2_children = service.get_level_children(
            target_level=2,
            parent_filters={'l1_id': self.customer_1.id, 'parent_node_id': f'n1_{self.customer_1.id}'},
        )
        self.assertEqual(len(l2_children), 2)  

        for child in l2_children:
            pc_id = child['id']
            totals_by_year = {t['year']: t for t in child['totals']}

            for year in [2024, 2025]:
                db_pc_year = SaleTransaction.objects.filter(
                    customer=self.customer_1,
                    product_class_id=pc_id,
                    sale_date__year=year,
                ).aggregate(net=Sum('net_amount'), profit=Sum('profit'))

                expected_net = float(db_pc_year['net'] or 0.0)
                expected_profit = float(db_pc_year['profit'] or 0.0)

                self.assertAlmostEqual(totals_by_year[year]['net'], expected_net, places=2)
                self.assertAlmostEqual(totals_by_year[year]['profit'], expected_profit, places=2)

        
        l3_children = service.get_level_children(
            target_level=3,
            parent_filters={
                'l1_id': self.customer_1.id,
                'l2_id': self.class_analgesic.id,
                'parent_node_id': f'n1_{self.customer_1.id}_{self.class_analgesic.id}',
            },
        )
        self.assertEqual(len(l3_children), 2)  

        for prod_child in l3_children:
            prod_id = prod_child['id']
            prod_totals = {t['year']: t for t in prod_child['totals']}

            for year in [2024, 2025]:
                db_prod_year = SaleTransaction.objects.filter(
                    customer=self.customer_1,
                    product_id=prod_id,
                    sale_date__year=year,
                ).aggregate(net=Sum('net_amount'), profit=Sum('profit'))

                expected_net = float(db_prod_year['net'] or 0.0)
                expected_profit = float(db_prod_year['profit'] or 0.0)

                self.assertAlmostEqual(prod_totals[year]['net'], expected_net, places=2)
                self.assertAlmostEqual(prod_totals[year]['profit'], expected_profit, places=2)

    def test_emitting_route_dimension_unaffected(self):
        """
        Validates that other dimensions (e.g. route_productclass_product) remain
        strictly bounded to the emitting route, NOT customer assignments.
        Route A should only reflect transactions where route=Route A (Customer 1 in 2025, Customer 2 in 2024).
        """
        tx_service = SaleTransactionsService(user=self.user)
        base_tx_qs = tx_service.read_transactions_by_allowed_routes()

        filter_set = YearlySaleBreakdownFilter(
            data={'dimension': 'route_productclass_product', 'route': [self.route_a.id]},
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_set.is_valid(), filter_set.errors)

        service = YearlySaleBreakdownService(
            queryset=filter_set.qs,
            dimension='route_productclass_product',
            user=self.user,
            cleaned_data=filter_set.form.cleaned_data,
        )

        l1_qs = service.get_level_1_queryset()
        l1_ids = [item['route_id'] for item in l1_qs]
        self.assertEqual(l1_ids, [self.route_a.id])

        items = service.get_level_1_items(l1_ids)
        self.assertEqual(len(items), 1)

        route_a_item = items[0]
        totals_by_year = {t['year']: t for t in route_a_item['totals']}

        
        for year in [2024, 2025]:
            db_route_year = SaleTransaction.objects.filter(
                route=self.route_a, sale_date__year=year
            ).aggregate(net=Sum('net_amount'), profit=Sum('profit'))

            expected_net = float(db_route_year['net'] or 0.0)
            expected_profit = float(db_route_year['profit'] or 0.0)

            self.assertAlmostEqual(totals_by_year[year]['net'], expected_net, places=2)
            self.assertAlmostEqual(totals_by_year[year]['profit'], expected_profit, places=2)

        
        self.assertAlmostEqual(totals_by_year[2024]['net'], 5000.0, places=2)
        
        self.assertAlmostEqual(totals_by_year[2025]['net'], 15000.0, places=2)

    def test_business_unit_and_region_filters_vs_db_raw(self):
        """
        Validates filtering by BusinessUnit and Region in customer_productclass_product
        against direct database aggregates of customers assigned to routes in those units.
        """
        today = timezone.localdate()
        tx_service = SaleTransactionsService(user=self.user)
        base_tx_qs = tx_service.read_transactions_by_allowed_customers()

        
        filter_bu = YearlySaleBreakdownFilter(
            data={'dimension': 'customer_productclass_product', 'business_unit': [self.bu_1.id]},
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_bu.is_valid(), filter_bu.errors)

        service_bu = YearlySaleBreakdownService(
            queryset=filter_bu.qs,
            dimension='customer_productclass_product',
            user=self.user,
            cleaned_data=filter_bu.form.cleaned_data,
        )

        expected_bu1_customers = list(
            CustomerAssignment.objects.filter(route__business_unit=self.bu_1)
            .filter(Q(end_date__isnull=True) | Q(end_date__gte=today))
            .values_list('customer_id', flat=True)
        )
        db_bu1_total = SaleTransaction.objects.filter(
            customer_id__in=expected_bu1_customers
        ).aggregate(net=Sum('net_amount'))

        l1_bu1_ids = [item['customer_id'] for item in service_bu.get_level_1_queryset()]
        self.assertEqual(l1_bu1_ids, expected_bu1_customers)

        items_bu1 = service_bu.get_level_1_items(l1_bu1_ids)
        service_bu1_net = sum(sum(t['net'] for t in item['totals']) for item in items_bu1)
        self.assertAlmostEqual(service_bu1_net, float(db_bu1_total['net']), places=2)

        
        filter_reg = YearlySaleBreakdownFilter(
            data={'dimension': 'customer_productclass_product', 'region': [self.region_north.id]},
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_reg.is_valid(), filter_reg.errors)

        service_reg = YearlySaleBreakdownService(
            queryset=filter_reg.qs,
            dimension='customer_productclass_product',
            user=self.user,
            cleaned_data=filter_reg.form.cleaned_data,
        )

        l1_reg_ids = [item['customer_id'] for item in service_reg.get_level_1_queryset()]
        self.assertEqual(l1_reg_ids, [self.customer_1.id])

        items_reg = service_reg.get_level_1_items(l1_reg_ids)
        service_reg_net = sum(sum(t['net'] for t in item['totals']) for item in items_reg)
        self.assertAlmostEqual(service_reg_net, float(db_bu1_total['net']), places=2)

    def test_csv_export_aggregates_and_route_assignment(self):
        """
        Validates CSV export generated by YearlySaleBreakdownExports:
        1. Net sales match direct DB aggregates.
        2. 'ID Ruta Asignada' and 'Ruta Asignada' display the current active assignment (Route A).
        """
        tx_service = SaleTransactionsService(user=self.user)
        base_tx_qs = tx_service.read_transactions_by_allowed_customers()

        filter_set = YearlySaleBreakdownFilter(
            data={'dimension': 'customer_productclass_product', 'route': [self.route_a.id]},
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_set.is_valid(), filter_set.errors)

        service = YearlySaleBreakdownService(
            queryset=filter_set.qs,
            dimension='customer_productclass_product',
            user=self.user,
            cleaned_data=filter_set.form.cleaned_data,
        )

        exporter = YearlySaleBreakdownExports(breakdown_service=service)
        csv_buffer = exporter.export_yearly_sale_breakdown_csv(is_seller=False)

        csv_content = csv_buffer.getvalue().decode('utf-8-sig')
        reader = list(csv.reader(io.StringIO(csv_content)))

        headers = reader[0]
        self.assertIn('ID Cliente', headers)
        self.assertIn('Cliente', headers)
        self.assertIn('ID Ruta Asignada', headers)
        self.assertIn('Ruta Asignada', headers)
        self.assertIn('Gerencia de Ruta', headers)

        
        idx_cid = headers.index('ID Cliente')
        idx_rid = headers.index('ID Ruta Asignada')
        idx_rname = headers.index('Ruta Asignada')
        idx_bu = headers.index('Gerencia de Ruta')

        data_rows = reader[1:]
        self.assertTrue(len(data_rows) > 0)

        
        for row in data_rows:
            self.assertEqual(row[idx_cid], self.customer_1.id)
            self.assertEqual(row[idx_rid], self.route_a.id)
            self.assertEqual(row[idx_rname].lower(), self.route_a.name.lower())
            self.assertEqual(row[idx_bu].lower(), self.bu_1.name.lower())

    def test_month_filter_aggregates_vs_db_raw(self):
        """
        Validates that filtering by specific months (e.g. March = '3')
        yields totals matching direct DB aggregates for those months only.
        """
        tx_service = SaleTransactionsService(user=self.user)
        base_tx_qs = tx_service.read_transactions_by_allowed_customers()

        
        filter_set = YearlySaleBreakdownFilter(
            data={
                'dimension': 'customer_productclass_product',
                'route': [self.route_a.id],
                'months': ['3'],
            },
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_set.is_valid(), filter_set.errors)

        service = YearlySaleBreakdownService(
            queryset=filter_set.qs,
            dimension='customer_productclass_product',
            user=self.user,
            cleaned_data=filter_set.form.cleaned_data,
        )

        db_march_total = SaleTransaction.objects.filter(
            customer=self.customer_1,
            sale_date__month=3,
        ).aggregate(net=Sum('net_amount'), profit=Sum('profit'))

        l1_ids = [item['customer_id'] for item in service.get_level_1_queryset()]
        items = service.get_level_1_items(l1_ids)
        self.assertEqual(len(items), 1)

        c1_item = items[0]
        totals_by_year = {t['year']: t for t in c1_item['totals']}

        self.assertIn(2024, totals_by_year)
        self.assertAlmostEqual(totals_by_year[2024]['net'], float(db_march_total['net']), places=2)
        self.assertAlmostEqual(totals_by_year[2024]['profit'], float(db_march_total['profit']), places=2)
        
        self.assertNotIn(2025, totals_by_year)

    def test_seller_export_aggregates_vs_db_raw(self):
        """
        Validates CSV export when is_seller=True:
        1. Ensures 'Margen <year>' header contains qualitative classification label for sellers.
        2. Net sales in columns match direct DB aggregates.
        """
        tx_service = SaleTransactionsService(user=self.user)
        base_tx_qs = tx_service.read_transactions_by_allowed_customers()

        filter_set = YearlySaleBreakdownFilter(
            data={'dimension': 'customer_productclass_product', 'route': [self.route_a.id]},
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_set.is_valid(), filter_set.errors)

        service = YearlySaleBreakdownService(
            queryset=filter_set.qs,
            dimension='customer_productclass_product',
            user=self.user,
            cleaned_data=filter_set.form.cleaned_data,
        )

        exporter = YearlySaleBreakdownExports(breakdown_service=service)
        csv_buffer = exporter.export_yearly_sale_breakdown_csv(is_seller=True)

        csv_content = csv_buffer.getvalue().decode('utf-8-sig')
        reader = list(csv.reader(io.StringIO(csv_content)))
        headers = reader[0]

        
        for y in service.sorted_years:
            self.assertIn(f'Margen {y}', headers)
            self.assertNotIn(f'Margen % {y}', headers)
            self.assertNotIn(f'Clasificación Margen {y}', headers)

        
        idx_net_2024 = headers.index('Venta Neta 2024')
        idx_margin_label_2024 = headers.index('Margen 2024')

        data_rows = reader[1:]
        total_csv_net_2024 = sum(float(r[idx_net_2024]) for r in data_rows)

        db_aggregates_2024 = SaleTransaction.objects.filter(
            customer=self.customer_1,
            sale_date__year=2024,
        ).aggregate(net=Sum('net_amount'), profit=Sum('profit'))

        self.assertAlmostEqual(total_csv_net_2024, float(db_aggregates_2024['net']), places=2)

        
        for r in data_rows:
            self.assertIn(r[idx_margin_label_2024], ['Excelente', 'Óptimo', 'Regular', 'Malo', 'Muy malo'])

    def test_productclass_customer_product_dimension_aggregates_vs_db_raw(self):
        """
        Validates productclass_customer_product dimension when filtering by Route A:
        1. Only considers actively assigned customers (Customer 1).
        2. Level 1 (ProductClass) totals reflect Customer 1's historical consumption across all years.
        3. Customer 2's sales (no longer assigned to Route A) are NOT included.
        4. Level 2 (Customer) only displays actively assigned customers.
        5. Compares all values directly against raw DB aggregates.
        """
        today = timezone.localdate()
        tx_service = SaleTransactionsService(user=self.user)
        base_tx_qs = tx_service.read_transactions_by_allowed_customers()

        filter_set = YearlySaleBreakdownFilter(
            data={'dimension': 'productclass_customer_product', 'route': [self.route_a.id]},
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_set.is_valid(), filter_set.errors)

        service = YearlySaleBreakdownService(
            queryset=filter_set.qs,
            dimension='productclass_customer_product',
            user=self.user,
            cleaned_data=filter_set.form.cleaned_data,
        )

        active_customer_ids = list(
            CustomerAssignment.objects.filter(route=self.route_a)
            .filter(Q(end_date__isnull=True) | Q(end_date__gte=today))
            .values_list('customer_id', flat=True)
        )
        self.assertEqual(active_customer_ids, [self.customer_1.id])

        
        l1_qs = service.get_level_1_queryset()
        l1_ids = [item['product_class_id'] for item in l1_qs]
        self.assertIn(self.class_analgesic.id, l1_ids)
        self.assertIn(self.class_antibiotic.id, l1_ids)

        items = service.get_level_1_items(l1_ids)
        items_map = {item['id']: item for item in items}

        
        
        analg_item = items_map[self.class_analgesic.id]
        analg_totals = {t['year']: t for t in analg_item['totals']}

        db_analg_2024 = SaleTransaction.objects.filter(
            customer_id__in=active_customer_ids,
            product_class=self.class_analgesic,
            sale_date__year=2024,
        ).aggregate(net=Sum('net_amount'), profit=Sum('profit'))

        db_analg_2025 = SaleTransaction.objects.filter(
            customer_id__in=active_customer_ids,
            product_class=self.class_analgesic,
            sale_date__year=2025,
        ).aggregate(net=Sum('net_amount'), profit=Sum('profit'))

        self.assertAlmostEqual(analg_totals[2024]['net'], float(db_analg_2024['net']), places=2)
        self.assertAlmostEqual(analg_totals[2024]['net'], 10000.0, places=2)  
        self.assertAlmostEqual(analg_totals[2025]['net'], float(db_analg_2025['net']), places=2)
        self.assertAlmostEqual(analg_totals[2025]['net'], 15000.0, places=2)

        
        l2_children = service.get_level_children(
            target_level=2,
            parent_filters={'l1_id': self.class_analgesic.id, 'parent_node_id': f'n1_{self.class_analgesic.id}'},
        )
        l2_customer_ids = [c['id'] for c in l2_children]
        self.assertEqual(l2_customer_ids, [self.customer_1.id])
        self.assertNotIn(self.customer_2.id, l2_customer_ids)

        c1_l2 = l2_children[0]
        c1_l2_totals = {t['year']: t for t in c1_l2['totals']}
        self.assertAlmostEqual(c1_l2_totals[2024]['net'], 10000.0, places=2)
        self.assertAlmostEqual(c1_l2_totals[2025]['net'], 15000.0, places=2)

        
        l3_children = service.get_level_children(
            target_level=3,
            parent_filters={
                'l1_id': self.class_analgesic.id,
                'l2_id': self.customer_1.id,
                'parent_node_id': f'n1_{self.class_analgesic.id}_{self.customer_1.id}',
            },
        )
        l3_product_ids = [p['id'] for p in l3_children]
        self.assertIn(self.prod_paracetamol.id, l3_product_ids)
        self.assertIn(self.prod_ibuprofen.id, l3_product_ids)

    def test_productclass_customer_product_csv_export(self):
        """
        Validates CSV export for productclass_customer_product dimension:
        1. Ensures assigned route columns correspond to active customer assignment.
        2. Customer 2 is excluded because it is no longer assigned to Route A.
        """
        tx_service = SaleTransactionsService(user=self.user)
        base_tx_qs = tx_service.read_transactions_by_allowed_customers()

        filter_set = YearlySaleBreakdownFilter(
            data={'dimension': 'productclass_customer_product', 'route': [self.route_a.id]},
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_set.is_valid(), filter_set.errors)

        service = YearlySaleBreakdownService(
            queryset=filter_set.qs,
            dimension='productclass_customer_product',
            user=self.user,
            cleaned_data=filter_set.form.cleaned_data,
        )

        exporter = YearlySaleBreakdownExports(breakdown_service=service)
        csv_buffer = exporter.export_yearly_sale_breakdown_csv(is_seller=False)

        csv_content = csv_buffer.getvalue().decode('utf-8-sig')
        reader = list(csv.reader(io.StringIO(csv_content)))
        headers = reader[0]

        self.assertIn('ID Clase de Producto', headers)
        self.assertIn('Clase de Producto', headers)
        self.assertIn('ID Cliente', headers)
        self.assertIn('Cliente', headers)
        self.assertIn('ID Ruta Asignada', headers)
        self.assertIn('Ruta Asignada', headers)
        self.assertIn('ID Producto', headers)

        idx_cid = headers.index('ID Cliente')
        idx_rid = headers.index('ID Ruta Asignada')

        data_rows = reader[1:]
        self.assertTrue(len(data_rows) > 0)

        for row in data_rows:
            self.assertEqual(row[idx_cid], self.customer_1.id)
            self.assertEqual(row[idx_rid], self.route_a.id)

    def test_product_customer_dimension_aggregates_vs_db_raw(self):
        """
        Validates product_customer dimension (Producto -> Cliente) when filtering by Route A:
        1. Level 1 (Product) only reflects active customers of Route A (Customer 1).
        2. Paracetamol excludes Customer 2's sales (ended assignment).
        3. Level 2 (Customer) only shows Customer 1 with its full consumption.
        4. CSV export correctly reports active route assignment.
        """
        today = timezone.localdate()
        tx_service = SaleTransactionsService(user=self.user)
        base_tx_qs = tx_service.read_transactions_by_allowed_customers()

        filter_set = YearlySaleBreakdownFilter(
            data={'dimension': 'product_customer', 'route': [self.route_a.id]},
            queryset=base_tx_qs,
        )
        self.assertTrue(filter_set.is_valid(), filter_set.errors)

        service = YearlySaleBreakdownService(
            queryset=filter_set.qs,
            dimension='product_customer',
            user=self.user,
            cleaned_data=filter_set.form.cleaned_data,
        )

        active_customer_ids = list(
            CustomerAssignment.objects.filter(route=self.route_a)
            .filter(Q(end_date__isnull=True) | Q(end_date__gte=today))
            .values_list('customer_id', flat=True)
        )
        self.assertEqual(active_customer_ids, [self.customer_1.id])

        
        l1_qs = service.get_level_1_queryset()
        l1_ids = [item['product_id'] for item in l1_qs]
        self.assertIn(self.prod_paracetamol.id, l1_ids)
        self.assertIn(self.prod_ibuprofen.id, l1_ids)
        self.assertIn(self.prod_amoxicillin.id, l1_ids)

        items = service.get_level_1_items(l1_ids)
        items_map = {item['id']: item for item in items}

        
        
        paracetamol_item = items_map[self.prod_paracetamol.id]
        paracetamol_totals = {t['year']: t for t in paracetamol_item['totals']}

        db_paracetamol_2024 = SaleTransaction.objects.filter(
            customer_id__in=active_customer_ids,
            product=self.prod_paracetamol,
            sale_date__year=2024,
        ).aggregate(net=Sum('net_amount'), profit=Sum('profit'))

        self.assertAlmostEqual(paracetamol_totals[2024]['net'], float(db_paracetamol_2024['net']), places=2)
        self.assertAlmostEqual(paracetamol_totals[2024]['net'], 10000.0, places=2)

        
        l2_children = service.get_level_children(
            target_level=2,
            parent_filters={'l1_id': self.prod_paracetamol.id, 'parent_node_id': f'n1_{self.prod_paracetamol.id}'},
        )
        l2_customer_ids = [c['id'] for c in l2_children]
        self.assertEqual(l2_customer_ids, [self.customer_1.id])
        self.assertNotIn(self.customer_2.id, l2_customer_ids)

        
        exporter = YearlySaleBreakdownExports(breakdown_service=service)
        csv_buffer = exporter.export_yearly_sale_breakdown_csv(is_seller=False)
        reader = list(csv.reader(io.StringIO(csv_buffer.getvalue().decode('utf-8-sig'))))
        headers = reader[0]

        self.assertIn('ID Producto', headers)
        self.assertIn('Producto', headers)
        self.assertIn('ID Cliente', headers)
        self.assertIn('Cliente', headers)
        self.assertIn('ID Ruta Asignada', headers)
        self.assertIn('Ruta Asignada', headers)

        idx_cid = headers.index('ID Cliente')
        idx_rid = headers.index('ID Ruta Asignada')

        for row in reader[1:]:
            self.assertEqual(row[idx_cid], self.customer_1.id)
            self.assertEqual(row[idx_rid], self.route_a.id)


class CustomerKpisAgreementsTestCase(TestCase):
    """
    test suite verifying that CustomerKpisService and CustomerProfileService
    dynamically calculate active customer agreements independent of filters,
    and verify template and export output.
    """

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            username='analytics_agreements_tester',
            email='tester@datall.local',
            is_superuser=True,
            is_staff=True,
        )
        cls.bu = BusinessUnit.objects.create(
            id='bu_kpis',
            name='Unidad KPIs',
            business_unit_type=BusinessUnit.BusinessUnitTypeChoices.UNIT,
        )
        cls.route_type = RouteType.objects.create(id='rt_kpi', name='Ruta KPI')
        cls.sale_channel = SaleChannel.objects.create(id='sc_kpi', name='Canal KPI')
        cls.route = Route.objects.create(
            id='R-KPI-1',
            name='Ruta KPI 1',
            business_unit=cls.bu,
            route_type=cls.route_type,
            sale_channel=cls.sale_channel,
            is_active=True,
        )
        cls.benefit = CommercialBenefit.objects.create(
            name='Descuento Especial',
            benefit_type='fixed_discount',
            cost=Decimal('500.00'),
            is_active=True,
        )
        cls.customer_type = CustomerType.objects.create(id='ct_kpi', name='Tipo KPI')
        today = timezone.localdate()
        cls.today = today

        cls.c_active = Customer.objects.create(
            id='C-ACT',
            name='Cliente Con Convenios Activos',
            registration_date=today - timedelta(days=365),
            customer_type=cls.customer_type,
        )
        cls.c_expired = Customer.objects.create(
            id='C-EXP',
            name='Cliente Solo Vencidos',
            registration_date=today - timedelta(days=365),
            customer_type=cls.customer_type,
        )
        cls.c_none = Customer.objects.create(
            id='C-NONE',
            name='Cliente Sin Convenios',
            registration_date=today - timedelta(days=365),
            customer_type=cls.customer_type,
        )

        
        CustomerAgreement.objects.create(
            customer=cls.c_active,
            route=cls.route,
            benefit=cls.benefit,
            start_date=today - timedelta(days=30),
            end_date=today + timedelta(days=30),
        )
        CustomerAgreement.objects.create(
            customer=cls.c_active,
            route=cls.route,
            benefit=cls.benefit,
            start_date=today - timedelta(days=10),
            end_date=None,
        )
        CustomerAgreement.objects.create(
            customer=cls.c_active,
            route=cls.route,
            benefit=cls.benefit,
            start_date=today - timedelta(days=90),
            end_date=today - timedelta(days=1),
        )
        CustomerAgreement.objects.create(
            customer=cls.c_active,
            route=cls.route,
            benefit=cls.benefit,
            start_date=today + timedelta(days=10),
            end_date=today + timedelta(days=60),
        )

        
        CustomerAgreement.objects.create(
            customer=cls.c_expired,
            route=cls.route,
            benefit=cls.benefit,
            start_date=today - timedelta(days=100),
            end_date=today - timedelta(days=10),
        )

    def _get_service(self, date_start=None, date_end=None):
        return CustomerKpisService(
            user=self.user,
            customers_qs=Customer.objects.filter(id__in=[self.c_active.id, self.c_expired.id, self.c_none.id]),
            transactions_qs=SaleTransaction.objects.none(),
            ars_qs=AccountsReceivable.objects.none(),
            date_start=date_start,
            date_end=date_end,
        )

    def test_get_active_agreements_map(self):
        service = self._get_service()
        agreements_map = service._get_active_agreements_map([
            self.c_active.id,
            self.c_expired.id,
            self.c_none.id,
        ])
        self.assertEqual(agreements_map.get(self.c_active.id), 2)
        self.assertIsNone(agreements_map.get(self.c_expired.id))
        self.assertIsNone(agreements_map.get(self.c_none.id))

    def test_read_customer_kpis_populates_active_agreements(self):
        service = self._get_service()
        results = service.read_customer_kpis()
        by_id = {c.id: c for c in results}

        self.assertEqual(by_id[self.c_active.id].active_agreements, 2)
        self.assertEqual(by_id[self.c_expired.id].active_agreements, 0)
        self.assertEqual(by_id[self.c_none.id].active_agreements, 0)

    def test_active_agreements_unaffected_by_date_filters(self):
        
        past_start = self.today - timedelta(days=60)
        past_end = self.today - timedelta(days=30)
        service = self._get_service(date_start=past_start, date_end=past_end)
        results = service.read_customer_kpis()
        by_id = {c.id: c for c in results}

        self.assertEqual(by_id[self.c_active.id].active_agreements, 2)
        self.assertEqual(by_id[self.c_expired.id].active_agreements, 0)
        self.assertEqual(by_id[self.c_none.id].active_agreements, 0)

    def test_customer_profile_service_sets_active_agreements(self):
        profile_service = CustomerProfileService(
            user=self.user,
            customer=self.c_active,
            customers_qs=Customer.objects.filter(id=self.c_active.id),
            transactions_qs=SaleTransaction.objects.none(),
            ars_qs=AccountsReceivable.objects.none(),
        )
        customer = profile_service.build_profile()
        self.assertEqual(customer.active_agreements, 2)

        expired_service = CustomerProfileService(
            user=self.user,
            customer=self.c_expired,
            customers_qs=Customer.objects.filter(id=self.c_expired.id),
            transactions_qs=SaleTransaction.objects.none(),
            ars_qs=AccountsReceivable.objects.none(),
        )
        expired_customer = expired_service.build_profile()
        self.assertEqual(expired_customer.active_agreements, 0)

    def test_customer_kpis_exports_excel_active_agreements(self):
        service = self._get_service()
        exporter = CustomerKpisExports(customer_kpis_service=service)
        excel_buffer = exporter.export_customer_kpis_report()
        wb = openpyxl.load_workbook(excel_buffer)
        ws = wb['Listado de Clientes']

        
        found = False
        for row in range(3, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == self.c_active.id:
                
                self.assertEqual(ws.cell(row=row, column=19).value, 2)
                found = True
                break
        self.assertTrue(found, "Customer C-ACT was not found in exported Excel")

    def test_template_rows_renders_active_agreements_link(self):
        service = self._get_service()
        results = service.read_customer_kpis()

        rendered = render_to_string(
            'analytics/customer_kpis/partials/_customer_kpis_rows.html',
            {'customers': results}
        )

        detail_url = reverse('customers:customer_detail_view', kwargs={'pk': self.c_active.id})
        self.assertIn(detail_url, rendered)
        self.assertIn('Ver convenios de Cliente Con Convenios Activos', rendered)
        self.assertIn('<span class="text-muted">0</span>', rendered)


class StockBreakdownTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(username='test_stock_user', password='password')

        cls.cat1 = ProductCategory.objects.create(id='CAT1', name='Farmacia')
        cls.cat2 = ProductCategory.objects.create(id='CAT2', name='Alimentos')

        cls.class1 = ProductClass.objects.create(id='CLS1', name='Antibióticos', product_category=cls.cat1)
        cls.class2 = ProductClass.objects.create(id='CLS2', name='Croquetas', product_category=cls.cat2)

        cls.prod1 = Product.objects.create(id='P1', name='Amoxicilina', product_class=cls.class1)
        cls.prod2 = Product.objects.create(id='P2', name='Cefalexina', product_class=cls.class1)
        cls.prod3 = Product.objects.create(id='P3', name='Diamond Adult', product_class=cls.class2)

        cls.wh_main = Warehouse.objects.create(
            id='wh_main',
            name='Almacén Central',
            warehouse_type=Warehouse.WarehouseTypeChoices.WAREHOUSE
        )
        cls.wh_sec = Warehouse.objects.create(
            id='wh_sec',
            name='Almacén Secundario',
            warehouse_type=Warehouse.WarehouseTypeChoices.WAREHOUSE
        )
        cls.wh_retail = Warehouse.objects.create(
            id='wh_ret',
            name='Tienda Norte',
            warehouse_type=Warehouse.WarehouseTypeChoices.RETAIL
        )

        Stock.objects.create(product=cls.prod1, warehouse=cls.wh_main, quantity=Decimal('100.00'), lot_number='L1')
        Stock.objects.create(product=cls.prod1, warehouse=cls.wh_sec, quantity=Decimal('50.00'), lot_number='L2')
        Stock.objects.create(product=cls.prod2, warehouse=cls.wh_main, quantity=Decimal('20.00'), lot_number='L3')
        Stock.objects.create(product=cls.prod3, warehouse=cls.wh_retail, quantity=Decimal('80.00'), lot_number='L4')

    def test_default_warehouses_resolution(self):
        service = StockBreakdownService(queryset=Stock.objects.all())
        wh_ids = [w.id for w in service.active_warehouses]
        self.assertIn('wh_main', wh_ids)
        self.assertIn('wh_sec', wh_ids)
        self.assertNotIn('wh_ret', wh_ids)

    def test_level_1_items_and_totals(self):
        service = StockBreakdownService(queryset=Stock.objects.all())
        l1_qs = service.get_level_1_queryset()
        self.assertEqual(l1_qs.count(), 1)  # Only CAT1 has stock in default warehouses

        top_ids = [i[service.l1_id_field] for i in l1_qs]
        items = service.get_level_1_items(top_ids)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]['id'], 'CAT1')
        self.assertEqual(items[0]['total_overall'], 170.0)

    def test_level_children(self):
        service = StockBreakdownService(queryset=Stock.objects.all())
        # Level 2 (classes) under CAT1
        l2_items = service.get_level_children(target_level=2, parent_filters={'l1_id': 'CAT1'})
        self.assertEqual(len(l2_items), 1)
        self.assertEqual(l2_items[0]['id'], 'CLS1')
        self.assertEqual(l2_items[0]['total_overall'], 170.0)

        # Level 3 (products) under CLS1
        l3_items = service.get_level_children(target_level=3, parent_filters={'l1_id': 'CAT1', 'l2_id': 'CLS1'})
        self.assertEqual(len(l3_items), 2)
        p1_item = next(i for i in l3_items if i['id'] == 'P1')
        self.assertEqual(p1_item['total_overall'], 150.0)
        self.assertTrue(p1_item['has_children'])
        self.assertTrue(p1_item['is_product'])

        # Level 4 (lots) under P1
        l4_items = service.get_level_children(target_level=4, parent_filters={'l1_id': 'CAT1', 'l2_id': 'CLS1', 'l3_id': 'P1'})
        self.assertEqual(len(l4_items), 2)
        lot_numbers = [item['lot_number'] for item in l4_items]
        self.assertIn('L1', lot_numbers)
        self.assertIn('L2', lot_numbers)
        self.assertTrue(all(item['is_lot'] for item in l4_items))
        self.assertTrue(all(not item['has_children'] for item in l4_items))

    def test_filter_selection_with_retail_warehouse(self):
        form_data = {
            'warehouse': [self.wh_retail.pk],
        }
        filter_set = StockBreakdownFilter(form_data, queryset=Stock.objects.all())
        self.assertTrue(filter_set.is_valid())
        service = StockBreakdownService(
            queryset=filter_set.qs,
            cleaned_data=filter_set.form.cleaned_data
        )
        l1_qs = service.get_level_1_queryset()
        self.assertEqual(l1_qs.count(), 1)
        self.assertEqual(l1_qs.first()[service.l1_id_field], 'CAT2')

    def test_excel_export_generation(self):
        service = StockBreakdownService(queryset=Stock.objects.all())
        exporter = StockBreakdownExports(breakdown_service=service)
        buf = exporter.export_stock_breakdown_excel()
        self.assertGreater(len(buf.getvalue()), 0)
        wb = openpyxl.load_workbook(buf)
        self.assertIn("Existencias por Almacén", wb.sheetnames)
        ws = wb["Existencias por Almacén"]
        # Verify lot rows are written
        found_lot = any("Lote:" in str(cell.value) for row in ws.iter_rows() for cell in row)
        self.assertTrue(found_lot)

    def test_kpis_calculation_and_buckets(self):
        from dateutil.relativedelta import relativedelta
        today = timezone.localdate()
        m0_start = today.replace(day=1)

        # Create additional test products with specific expiration dates
        prod_neg = Product.objects.create(id='P_NEG', name='Prod Negativo', product_class=self.class1)
        prod_zero = Product.objects.create(id='P_ZERO', name='Prod Cero', product_class=self.class1)
        prod_exp = Product.objects.create(id='P_EXP', name='Prod Expirando', product_class=self.class1)

        # Negative stock should be counted as having stock
        Stock.objects.create(product=prod_neg, warehouse=self.wh_main, quantity=Decimal('-5.00'), lot_number='L_NEG')
        # Zero stock should NOT be counted as having stock
        Stock.objects.create(product=prod_zero, warehouse=self.wh_main, quantity=Decimal('0.00'), lot_number='L_ZERO')

        # Expired (<= today, in previous month)
        Stock.objects.create(
            product=prod_exp,
            warehouse=self.wh_main,
            quantity=Decimal('10.00'),
            lot_number='L_EXPIRED',
            expiration_date=m0_start - timedelta(days=5)
        )
        # Expires in 0-2 months (month 0 or 1)
        Stock.objects.create(
            product=prod_exp,
            warehouse=self.wh_sec,
            quantity=Decimal('15.00'),
            lot_number='L_M01',
            expiration_date=m0_start + relativedelta(days=12)
        )
        # Expires in 3-4 months (month 2 or 3)
        Stock.objects.create(
            product=prod_exp,
            warehouse=self.wh_main,
            quantity=Decimal('25.00'),
            lot_number='L_M23',
            expiration_date=m0_start + relativedelta(months=2, days=5)
        )
        # Expires in 5-6 months (month 4 or 5)
        Stock.objects.create(
            product=prod_exp,
            warehouse=self.wh_main,
            quantity=Decimal('35.00'),
            lot_number='L_M45',
            expiration_date=m0_start + relativedelta(months=4, days=5)
        )
        # Expires +6 months (month 6+)
        Stock.objects.create(
            product=prod_exp,
            warehouse=self.wh_main,
            quantity=Decimal('45.00'),
            lot_number='L_M6P',
            expiration_date=m0_start + relativedelta(months=7, days=5)
        )

        service = StockBreakdownService(queryset=Stock.objects.all())
        kpis = service.get_kpis()

        # Prod1 (150), Prod2 (20), Prod_neg (-5), Prod_exp (130) -> 4 products with stock != 0. Prod_zero is excluded.
        self.assertEqual(kpis['products_with_stock'], 4)
        # Total stock = 100 + 50 + 20 - 5 + 0 + 10 + 15 + 25 + 35 + 45 = 295.00
        self.assertEqual(kpis['total_stock'], 295.0)
        # Expired units <= today = 10.0
        self.assertEqual(kpis['expired_units'], 10.0)
        # Expired products = 1 (prod_exp)
        self.assertEqual(kpis['expired_products'], 1)
        # Expiration buckets
        self.assertEqual(kpis['expires_0_2_months'], 15.0)
        self.assertEqual(kpis['expires_3_4_months'], 25.0)
        self.assertEqual(kpis['expires_5_6_months'], 35.0)
        self.assertEqual(kpis['expires_6_plus_months'], 45.0)

    def test_stock_breakdown_views_and_child_rendering(self):
        self.client.force_login(self.user)
        url = reverse('analytics:stock_breakdown_view')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn('kpis', response.context)
        self.assertGreater(response.context['kpis']['total_stock'], 0)

        # Lazy load children at level 4 (lots)
        child_url = reverse('analytics:stock_breakdown_children_view')
        child_resp = self.client.get(child_url, {
            'level': 4,
            'l1_id': 'CAT1',
            'l2_id': 'CLS1',
            'l3_id': 'P1',
            'parent_node_id': 'n1_CAT1_CLS1_P1'
        })
        self.assertEqual(child_resp.status_code, 200)
        content = child_resp.content.decode('utf-8')
        self.assertIn('Lote: L1', content)
        self.assertIn('Lote: L2', content)






